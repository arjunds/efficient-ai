#!/usr/bin/env python3
"""Build PI_BRIEF.docx (figures embedded) from PI_BRIEF.md. Minimal markdown subset:
headings, paragraphs, bullets, numbered lists, pipe tables, images, fenced code,
inline **bold**, *italic*, `code`. Run from GPU_power/ with pydeps on PYTHONPATH."""
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC, DST = "PI_BRIEF.md", "PI_BRIEF.docx"
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)

TOK = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(p, text, italic=False):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2].replace("`", "")); r.bold = True; r.italic = italic
        elif part.startswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.italic = italic
        elif part.startswith("*") and len(part) > 2:
            r = p.add_run(part[1:-1]); r.italic = True
        else:
            r = p.add_run(part); r.italic = italic


lines = open(SRC, encoding="utf-8").read().splitlines()
i = 0; n_img = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1; continue
    if ln.startswith("```"):
        i += 1; code = []
        while i < len(lines) and not lines[i].startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph(); r = p.add_run("\n".join(code))
        r.font.name = "Consolas"; r.font.size = Pt(9.5)
        i += 1; continue
    m = re.match(r"^(#{1,4})\s+(.*)", ln)
    if m:
        doc.add_heading(m.group(2), level=min(len(m.group(1)) - 1, 3) or 0)
        i += 1; continue
    m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", ln)
    if m:
        doc.add_picture(m.group(2), width=Inches(6.3))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        n_img += 1; i += 1; continue
    if ln.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
        for ri, row in enumerate(rows):
            for ci, c in enumerate(row):
                cell = t.cell(ri, ci); cell.text = ""
                add_inline(cell.paragraphs[0], f"**{c}**" if ri == 0 and not c.startswith("**") else c)
        doc.add_paragraph()
        continue
    m = re.match(r"^(\s*)[-*]\s+(.*)", ln)
    if m and not ln.startswith("*Figure") and not ln.startswith("*Status"):
        p = doc.add_paragraph(style="List Bullet" if not m.group(1) else "List Bullet 2")
        add_inline(p, m.group(2)); i += 1; continue
    m = re.match(r"^\d+\.\s+(.*)", ln)
    if m:
        p = doc.add_paragraph(style="List Number"); add_inline(p, m.group(1)); i += 1; continue
    p = doc.add_paragraph()
    if ln.startswith("*") and ln.endswith("*") and not ln.startswith("**"):
        add_inline(p, ln[1:-1], italic=True)
        for r in p.runs:
            r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    else:
        add_inline(p, ln)
    i += 1

doc.save(DST)
print("images embedded:", n_img)
